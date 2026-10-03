from pathlib import Path
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from modfilegen.utils.stics_cultivar_file import SticsCultivarFile
from modfilegen.utils.stics_txt_to_xml import (
    convert_plant_txt_to_xml,
    convert_soil_txt_to_xml,
    convert_initialization_txt_to_xml,
    convert_technical_txt_to_xml,
    convert_station_txt_to_xml,
    convert_usm_directory_to_xml,
    convert_general_parameters_txt_to_xml,
)


def _write(path: Path, text: str) -> Path:
    path.write_text(text, encoding="utf-8")
    return path


def _template(path: Path) -> Path:
    return _write(
        path,
        """<?xml version="1.0" encoding="UTF-8"?>
<fichierplt version="11.0">
  <formalisme nom="species">
    <param format="character" nom="codeplante">template</param>
    <option choix="1" nom="option" nomParam="codeoption">
      <choix code="1" nom="one"/>
      <choix code="2" nom="two"/>
    </option>
    <tv nb_varietes="1" nom="genotypes">
      <variete nom="template-variety">
        <formalismev nom="development">
          <param format="real" nom="stlevamf">100</param>
          <param format="real" nom="hautmax">2.5</param>
        </formalismev>
      </variete>
    </tv>
  </formalisme>
</fichierplt>
""",
    )


def test_convert_plant_txt_to_xml_replaces_values_and_clones_varieties(tmp_path):
    source = _write(
        tmp_path / "plant.txt",
        """codeplante
mai
codeoption
2
codevar
early
stlevamf
73
hautmax
2.2
codevar
late
stlevamf
140
hautmax
3.1
""",
    )
    destination = tmp_path / "plant.xml"

    result = convert_plant_txt_to_xml(
        source,
        _template(tmp_path / "template.xml"),
        destination,
    )

    generated_text = destination.read_text(encoding="utf-8")
    assert generated_text.startswith(
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
    )
    assert " />" not in generated_text
    root = ET.parse(destination).getroot()
    assert root.find(".//param[@nom='codeplante']").text == "mai"
    assert root.find(".//option[@nomParam='codeoption']").get("choix") == "2"
    table = root.find(".//tv")
    assert table.get("nb_varietes") == "2"
    varieties = table.findall("variete")
    assert [variety.get("nom") for variety in varieties] == ["early", "late"]
    assert [
        variety.find(".//param[@nom='stlevamf']").text for variety in varieties
    ] == ["73", "140"]
    assert [
        variety.find(".//param[@nom='hautmax']").text for variety in varieties
    ] == ["2.2", "3.1"]
    assert result.cultivars == 2
    assert result.species_parameters_updated == 2
    assert result.cultivar_parameters_updated == 4
    assert result.template_species_defaults == ()
    assert result.template_cultivar_defaults == ()


def test_template_only_parameter_requires_explicit_non_strict_mode(tmp_path):
    source = _write(
        tmp_path / "plant.txt",
        """codeplante
mai
codevar
only
stlevamf
73
""",
    )
    template = _template(tmp_path / "template.xml")

    with pytest.raises(ValueError, match="parameters from XML template absent from TXT"):
        convert_plant_txt_to_xml(source, template, tmp_path / "strict.xml")

    result = convert_plant_txt_to_xml(
        source,
        template,
        tmp_path / "defaults.xml",
        strict=False,
    )
    assert result.template_species_defaults == ("codeoption",)
    assert result.template_cultivar_defaults == ("hautmax",)


def test_txt_parameter_absent_from_template_is_always_an_error(tmp_path):
    source = _write(
        tmp_path / "plant.txt",
        """codeplante
mai
unknown
1
codevar
only
stlevamf
73
hautmax
2.5
""",
    )

    with pytest.raises(ValueError, match="unknown"):
        convert_plant_txt_to_xml(
            source,
            _template(tmp_path / "template.xml"),
            tmp_path / "output.xml",
            strict=False,
        )


def test_empty_template_has_clear_error(tmp_path):
    source = _write(
        tmp_path / "plant.txt",
        """codeplante
mai
codevar
only
stlevamf
73
hautmax
2.5
""",
    )
    template = _write(tmp_path / "empty.xml", "")

    with pytest.raises(ValueError, match="XML template is empty"):
        convert_plant_txt_to_xml(source, template, tmp_path / "output.xml")


@pytest.mark.parametrize("source_name", ["maiplt1_c850.txt", "maiplt1_c900.txt"])
def test_real_plant_fixtures_are_fully_transferred(source_name, tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "plant"
    source_path = fixture_dir / source_name
    destination = tmp_path / f"{source_path.stem}.xml"
    source = SticsCultivarFile.read(source_path)

    result = convert_plant_txt_to_xml(
        source_path,
        fixture_dir / "plant_template.xml",
        destination,
    )

    root = ET.parse(destination).getroot()
    varieties = list(root.iter("variete"))
    assert result.species_parameters_updated == 252
    assert result.cultivar_parameters_updated == 52
    assert len(varieties) == len(source.cultivars) == 1
    assert root.find(".//tv").get("nb_varietes") == "1"

    cultivar_descendants = {id(element) for variety in varieties for element in variety.iter()}
    species_values = {}
    for element in root.iter():
        if id(element) in cultivar_descendants:
            continue
        if element.tag == "param" and element.get("nom"):
            species_values[element.get("nom")] = element.text
        elif element.tag == "option" and element.get("nomParam"):
            species_values[element.get("nomParam")] = element.get("choix")
    assert species_values == dict(source.species_parameters)

    for variety, (cultivar_name, expected_values) in zip(
        varieties, source.cultivars.items()
    ):
        assert variety.get("nom") == cultivar_name
        actual_values = {
            element.get("nom"): element.text for element in variety.iter("param")
        }
        assert actual_values == dict(expected_values)


def test_real_soil_files_are_combined_and_fully_transferred(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "soil"
    sources = [fixture_dir / "AUS_A_MZ_88.sol", fixture_dir / "NIG_C_MZ_95.sol"]
    destination = tmp_path / "sols.xml"

    result = convert_soil_txt_to_xml(
        sources, fixture_dir / "template_sols.xml", destination
    )

    text = destination.read_text(encoding="utf-8")
    assert text.startswith('<?xml version="1.0" encoding="UTF-8"?>\n')
    assert "standalone" not in text.splitlines()[0]
    assert " />" not in text
    assert "solcarotte" not in text
    root = ET.parse(destination).getroot()
    soils = root.findall("sol")
    assert [soil.get("nom") for soil in soils] == ["AUS_A_MZ_88", "NIG_C_MZ_95"]
    assert result.soils == 2
    assert result.soil_parameters_updated == 66
    assert result.layer_parameters_updated == 80

    aus, nig = soils
    assert aus.find(".//param[@nom='argi']").text == "33.0"
    assert aus.find(".//param[@nom='CsurNsol']").text == "12.5000"
    assert aus.find(".//option[@nomParam='codrainage']").get("choix") == "2"
    assert nig.find(".//param[@nom='argi']").text == "5.5"
    assert nig.find(".//param[@nom='obstarac']").text == "115.0000"
    for soil, epc, hccf in ((aus, "40.00", "21.02"), (nig, "23.00", "17.91")):
        layers = soil.findall("tableau")
        assert len(layers) == 5
        assert [layer.get("nom") for layer in layers] == [f"layer {i}" for i in range(1, 6)]
        assert all(layer.find("colonne[@nom='epc']").text == epc for layer in layers)
        assert all(layer.find("colonne[@nom='HCCF']").text == hccf for layer in layers)


def test_soil_file_requires_three_headers_and_five_layers(tmp_path):
    source = _write(tmp_path / "short.sol", "1 Sol 1\n")
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "soil"
    with pytest.raises(ValueError, match="expected 8 non-empty lines"):
        convert_soil_txt_to_xml(
            [source], fixture_dir / "template_sols.xml", tmp_path / "sols.xml"
        )


def test_real_initialization_file_is_fully_transferred(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "initialization"
    destination = tmp_path / "mais_ini.xml"

    result = convert_initialization_txt_to_xml(
        fixture_dir / "ficini.txt",
        fixture_dir / "template_mais_ini.xml",
        destination,
    )

    text = destination.read_text(encoding="utf-8")
    assert text.startswith('<?xml version="1.0" encoding="UTF-8"?>\n')
    assert "standalone" not in text.splitlines()[0]
    assert " />" not in text
    root = ET.parse(destination).getroot()
    assert root.find("nbplantes").text == "1"
    plant = root.findall("plante")[0]
    assert plant.find("stade0").text == "snu"
    assert plant.find("option[@nomParam='code_acti_reserve']").get("choix") == "2"
    assert [h.text for h in plant.findall("densinitial/horizon")] == ["0.0"] * 5
    assert [h.text for h in root.findall("sol/Hinitf/horizon")] == ["12.7389"] * 5
    assert [h.text for h in root.findall("sol/NO3initf/horizon")] == ["0.0"] * 5
    assert [h.text for h in root.findall("sol/NH4initf/horizon")] == ["0.0"] * 5
    assert [root.find(f"snow/{name}").text for name in ("Sdepth0", "Sdry0", "Swet0", "ps0")] == ["0"] * 4
    assert result.plants == 1
    assert result.plant_values_updated == 17
    assert result.soil_values_updated == 15
    assert result.snow_values_updated == 4


def test_real_technical_file_transfers_parameters_and_optional_interventions(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "tec"
    destination = tmp_path / "Mais_tec.xml"

    result = convert_technical_txt_to_xml(
        fixture_dir / "fictec1.txt",
        fixture_dir / "template_Mais_tec.xml",
        destination,
    )

    text = destination.read_text(encoding="utf-8")
    assert text.startswith(
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
    )
    assert " />" not in text
    root = ET.parse(destination).getroot()
    assert root.find(".//param[@nom='iplt0']").text == "33"
    assert root.find(".//param[@nom='densitesem']").text == "7.00"
    assert root.find(".//option[@nomParam='codestade']").get("choix") == "2"
    assert root.find(".//param[@nom='irecbutoir']").text == "171"
    tables = list(root.iter("ta"))
    assert [table.get("nb_interventions") for table in tables] == [
        "0", "0", "0", "1", "0", "0", "0"
    ]
    assert [len(table.findall("intervention")) for table in tables] == [
        0, 0, 0, 1, 0, 0, 0
    ]
    fertilization = tables[3].findall("intervention/colonne")
    assert [(column.get("nom"), column.text) for column in fertilization] == [
        ("julapN_or_sum_upvt", "32"),
        ("absolute_value/%", "0.0"),
        ("engrais", "3"),
    ]
    assert result.parameters_updated == 114
    assert result.intervention_tables == 7
    assert result.interventions == 1


def test_technical_file_accepts_repeated_irrigation_headers(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "tec"
    source_text = (fixture_dir / "fictec1.txt").read_text(encoding="utf-8")
    source_text = source_text.replace(
        "codedateappH2O\n2\nnbinterventions\n0\ncodlocirrig",
        "codedateappH2O\n2\nnbinterventions\n2\n"
        "julapI_or_sum_upvt amount\n33 20.0\n"
        "julapI_or_sum_upvt amount\n68 50.0\ncodlocirrig",
    )
    source = _write(tmp_path / "fictec1.txt", source_text)
    destination = tmp_path / "Mais_tec.xml"

    result = convert_technical_txt_to_xml(
        source,
        fixture_dir / "template_Mais_tec.xml",
        destination,
    )

    irrigation = list(ET.parse(destination).getroot().iter("ta"))[2]
    assert irrigation.get("nb_interventions") == "2"
    assert [
        [(column.get("nom"), column.text) for column in row.findall("colonne")]
        for row in irrigation.findall("intervention")
    ] == [
        [("julapI_or_sum_upvt", "33"), ("amount", "20.0")],
        [("julapI_or_sum_upvt", "68"), ("amount", "50.0")],
    ]
    assert result.interventions == 3


def test_real_station_file_is_fully_transferred(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "station"
    destination = tmp_path / "statio5j_sta.xml"

    result = convert_station_txt_to_xml(
        fixture_dir / "station.txt",
        fixture_dir / "template_statio5j_sta.xml",
        destination,
    )

    text = destination.read_text(encoding="utf-8")
    assert text.startswith(
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
    )
    assert " />" not in text
    root = ET.parse(destination).getroot()
    assert root.find(".//param[@nom='latitude']").text == "-14.4600000"
    assert root.find(".//option[@nomParam='codeetp']").get("choix") == "2"
    assert root.find(".//param[@nom='altisimul']").text == "108.00000"
    assert root.find(".//option[@nomParam='codernet']").get("choix") == "2"
    assert root.find(".//option[@nomParam='codemodlsnow']").get("choix") == "3"
    assert root.find(".//param[@nom='tminseuil']").text == "-0.5"
    assert result.parameters_updated == 44
    assert result.template_defaults == ()


def test_station_missing_parameter_requires_explicit_non_strict_mode(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "station"
    source = _write(tmp_path / "station.txt", "zr\n2.5\n")
    template = fixture_dir / "template_statio5j_sta.xml"

    with pytest.raises(ValueError, match="parameters from XML template absent from TXT"):
        convert_station_txt_to_xml(source, template, tmp_path / "strict.xml")
    result = convert_station_txt_to_xml(
        source, template, tmp_path / "defaults.xml", strict=False
    )
    assert "latitude" in result.template_defaults


def test_usm_directory_generates_usms_associated_xml_and_yearly_climates(tmp_path):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml"
    output = tmp_path / "generated"

    result = convert_usm_directory_to_xml(
        fixture_dir / "usms",
        fixture_dir / "usms.xml",
        output,
        plant_template_xml=fixture_dir / "plant" / "plant_template.xml",
        initialization_template_xml=fixture_dir / "initialization" / "template_mais_ini.xml",
        soil_template_xml=fixture_dir / "soil" / "template_sols.xml",
        technical_template_xml=fixture_dir / "tec" / "template_Mais_tec.xml",
        station_template_xml=fixture_dir / "station" / "template_statio5j_sta.xml",
        master_input_db=fixture_dir / "usms" / "MasterInput.db",
        summary_template_xlsx=fixture_dir / "usms" / "QN_STICS.xlsx",
        general_parameters_xml=(
            fixture_dir / "general_parameters" / "generated_xml" / "param_gen.xml"
        ),
        new_form_parameters_xml=(
            fixture_dir / "general_parameters" / "generated_xml" / "param_newform.xml"
        ),
    )

    root = ET.parse(output / "usms.xml").getroot()
    usms = root.findall("usm")
    expected_names = sorted(
        path.name for path in (fixture_dir / "usms").iterdir() if path.is_dir()
    )
    assert [usm.get("nom") for usm in usms] == expected_names
    assert len(usms) == result.usms == 8
    for usm in usms:
        assert usm.find("nomsol").text == usm.get("nom")
        for field in ("finit", "fstation", "fclim1", "fclim2"):
            assert (output / usm.find(field).text).is_file()
        plant = usm.find("plante[@dominance='1']")
        assert (output / "plant" / plant.find("fplt").text).is_file()
        assert (output / plant.find("ftec").text).is_file()
        assert plant.find("fobs").text == f"{usm.get('nom')}.obs"
        assert (output / plant.find("fobs").text).is_file()

    soils = ET.parse(output / "sols.xml").getroot().findall("sol")
    assert [soil.get("nom") for soil in soils] == expected_names
    assert result.soils == 8
    assert result.climate_files == 8
    assert result.associated_xml_files == 28
    assert sorted(path.name for path in (output / "plant").glob("*.xml")) == [
        "Grinkan_plt.xml", "var_1600_plt.xml", "var_2100_plt.xml",
    ]
    shared_1600 = {
        usm.find("plante[@dominance='1']/fplt").text
        for usm in usms
        if usm.get("nom").startswith("AUS_B_")
    }
    assert shared_1600 == {"var_1600_plt.xml"}
    assert result.observation_files == 8
    assert result.summary_workbook == output / "QN_STICS.xlsx"
    assert result.summary_workbook.is_file()
    workbook = load_workbook(result.summary_workbook, read_only=True)
    assert "rapport" in workbook.sheetnames
    assert "dailyOutput" in workbook.sheetnames
    assert [cell.value for cell in workbook["rapport"][1]] == [
        "N°", "Model", "IdSim", "Texte", "iplts", "ilevs", "iflos", "imats",
        "masec(n)", "mafruit", "chargefruit", "laimax", "Qles", "QNapp",
        "QNplante", "ces", "cep", "SeasonOrder",
    ]
    assert workbook["rapport"].max_row == 144
    assert workbook["dailyOutput"].max_row > 1
    climate_1985 = (output / "cliAUS_Aj.1985").read_text(encoding="utf-8")
    assert climate_1985.endswith("\n")
    assert {line.split()[1] for line in climate_1985.splitlines()} == {"1985"}
    assert len(climate_1985.splitlines()) == 365

    aus_t1 = next(usm for usm in usms if usm.get("nom") == "AUS_A_MZ_T1_86")
    assert aus_t1.find("datedebut").text == "363"
    assert aus_t1.find("datefin").text == "533"
    assert aus_t1.find("fclim1").text == "cliAUS_Aj.1985"
    assert aus_t1.find("fclim2").text == "cliAUS_Aj.1986"
    assert (output / "AUS_A_MZ_88.obs").read_text(encoding="utf-8") == (
        "ian;mo;jo;jul;iplts;imats\n"
        "1988;5;19;140;33;140\n"
    )
    assert (output / "AUS_A_MZ_T1_86.obs").read_text(encoding="utf-8") == (
        "ian;mo;jo;jul;iplts;imats\n"
        "1986;5;17;137;394;502\n"
    )

    workbook = load_workbook(result.summary_workbook, data_only=True)
    assert workbook.sheetnames == [
        "USMs", "Ini", "Soils", "Tec", "Station", "Obs", "Weather", "Plant",
        "GeneralParameters", "NewFormParameters", "rapport", "dailyOutput",
        "MaturityComparison",
    ]
    comparison = workbook["MaturityComparison"]
    assert comparison.max_row == 144
    assert len(comparison._charts) == 1
    comparison_headers = [cell.value for cell in comparison[1]]
    assert comparison_headers == [
        "IdSim", "idPoint", "ObservedHarvestDate", "SimulatedMaturityDate",
        "ErrorDays", "SeasonOrder",
    ]
    assert workbook["USMs"].max_row == 9
    usm_sheet = workbook["USMs"]
    usm_name_column = [cell.value for cell in usm_sheet[1]].index("usm_name") + 1
    assert [
        usm_sheet.cell(row=row, column=usm_name_column).value
        for row in range(2, usm_sheet.max_row + 1)
    ] == expected_names
    assert workbook["Ini"].max_row == 9
    assert workbook["Soils"].max_row == 9
    assert workbook["Tec"].max_row == 9
    assert workbook["Station"].max_row == 9
    assert workbook["Obs"].max_row == 9
    assert workbook["Obs"]["G3"].value == 502
    assert workbook["Weather"].max_row == 2923
    assert workbook["Plant"].max_row > 100
    assert workbook["GeneralParameters"].max_row == 522
    assert workbook["NewFormParameters"].max_row == 26
    general_rows = list(workbook["GeneralParameters"].iter_rows(min_row=2, values_only=True))
    croco_rows = [row for row in general_rows if row[3] == "CroCo"]
    assert len(croco_rows) == 21
    assert [row[4] for row in croco_rows] == list(range(1, 22))
    new_form_rows = list(workbook["NewFormParameters"].iter_rows(min_row=2, values_only=True))
    assert next(row for row in new_form_rows if row[3] == "code_humirac")[5] == 1


@pytest.mark.parametrize(
    ("source_name", "template_name", "expected_count"),
    [
        ("tempopar.sti", "param_gen.xml", 521),
        ("tempoparv6.sti", "param_newform.xml", 25),
    ],
)
def test_general_parameters_preserve_every_ordered_occurrence(
    source_name, template_name, expected_count, tmp_path
):
    fixture_dir = Path(__file__).parent / "stics_txt_to_xml" / "general_parameters"
    destination = tmp_path / template_name
    result = convert_general_parameters_txt_to_xml(
        fixture_dir / source_name, fixture_dir / template_name, destination
    )

    source_lines = [
        line.strip()
        for line in (fixture_dir / source_name).read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    expected = {}
    for index in range(0, len(source_lines), 2):
        name = "code_humirac" if source_lines[index] == "humirac" else source_lines[index]
        expected.setdefault(name, []).append(source_lines[index + 1])
    actual = {}
    root = ET.parse(destination).getroot()
    for element in root.iter():
        name = element.get("nom") if element.tag == "param" else (
            element.get("nomParam") if element.tag == "option" else None
        )
        if name:
            value = element.get("choix") if element.tag == "option" else element.text
            actual.setdefault(name, []).append(value)

    assert actual == expected
    assert result.values_updated == expected_count
    text = destination.read_text(encoding="utf-8")
    assert text.startswith(
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
    )
    assert " />" not in text
    if source_name == "tempopar.sti":
        assert len(actual["CroCo"]) == 21
        assert actual["CroCo"] == expected["CroCo"]
